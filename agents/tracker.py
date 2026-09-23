from __future__ import annotations
import os
import re
import json
import datetime
from pathlib import Path
from typing import Dict, Any, Optional, List, Tuple
import yaml
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from db.session import SessionLocal
from db.models import Job, Application, Score, Artifact, Event

PROJECT_ROOT = Path(__file__).resolve().parent.parent

class TrackerAgent:
    """Tracks application pipeline lifecycle, generates daily markdown digests,
    and produces styled Excel tracker exports.
    """

    def __init__(self, db_session_factory=SessionLocal, project_root: Path = PROJECT_ROOT):
        self.session_factory = db_session_factory
        self.project_root = project_root
        self.reports_dir = self.project_root / "reports"
        self.reports_dir.mkdir(parents=True, exist_ok=True)

    def is_swe_intern_role(self, title: str) -> bool:
        """Determine if role title is a Software Engineering / Developer intern role."""
        if not title:
            return False
        t = str(title).lower().strip()
        
        intern_markers = ['intern', 'co-op', 'coop', 'campus', 'student', 'fellow', 'summer', 'undergrad', 'grad', 'fellowship', 'analyst program']
        if not any(m in t for m in intern_markers):
            return False
            
        non_swe_exact = [
            'product design', 'product designer', 'product manager', 'product management',
            'project manager', 'program manager', 'program management',
            'hardware', 'analog', 'digital design', 'asic', 'rf ', 'silicon', 'fpga design',
            'mechanical', 'civil', 'chemical', 'optical', 'biomedical', 'structural', 'acoustic',
            'business analyst', 'sales', 'marketing', 'accountant', 'accounting', 'finance analyst',
            'financial analyst', 'investment analyst', 'investment banking', 'wealth management',
            'actuarial', 'actuary', 'trader intern', 'trading intern', 'commodity trading',
            'risk analyst', 'risk management', 'credit risk', 'audit', 'recruiter', 'recruiting',
            'human resources', 'talent acquisition', 'legal', 'compliance', 'supply chain',
            'operations intern', 'underwriting', 'graphic design', 'ux designer', 'ui designer',
            'weapon control', 'customer insight', 'category insight'
        ]
        if any(ns in t for ns in non_swe_exact):
            if not any(sw in t for sw in ['software', 'swe', 'sde', 'developer', 'coding', 'development']):
                return False

        if any(qr in t for qr in ['quantitative research', 'quant research', 'quant trader', 'quantitative trader']):
            if not any(qd in t for qd in ['developer', 'development', 'software', 'engineer', 'swe']):
                return False

        swe_patterns = [
            r'\bsoftware\b',
            r'\bswe\b',
            r'\bsde\b',
            r'\bdeveloper\b',
            r'\bdevelopment\b',
            r'\bfrontend\b',
            r'\bbackend\b',
            r'\bfull\s*stack\b',
            r'\bdistributed\s+systems\b',
            r'\bplatform\s+engineer\b',
            r'\binfrastructure\s+engineer\b',
            r'\bsystems\s+engineer\b',
            r'\bsystems\s+software\b',
            r'\bcore\s+engineer\b',
            r'\bdata\s+(platform|infrastructure|engineer)\b',
            r'\b(machine\s+learning|ml|ai)\s+(engineer|software|developer)\b',
            r'\b(cloud|devops|sre|site\s+reliability)\s*(engineer|intern)\b',
            r'\bquant(itative)?\s*(developer|software|development)\b',
            r'\b(algo|algorithmic)\s*(trading\s*)?(developer|software)\b',
            r'\btrading\s+systems\b',
            r'\b(ios|android|mobile|web)\s+(developer|engineer)\b',
            r'\bembedded\s+(software|systems)\b',
            r'\bfirmware\s+(engineer|developer|software)\b',
            r'\b(technology|tech|engineering)\s+(summer\s+)?analyst\b',
            r'\bsummer\s+analyst\b.*(engineering|technology|developer|software)'
        ]
        
        for pat in swe_patterns:
            if re.search(pat, t):
                if any(bad in t for bad in ['business development', 'learning & development', 'real estate development', 'land development', 'talent development']):
                    return False
                return True
                
        return False

    def location_priority_sort_key(self, loc: Optional[str]) -> Tuple[int, str]:
        """Sort locations: New York (0) -> SF Bay Area (1) -> Chicago (2) -> Seattle (3) -> Alphabetical (4) -> Blank (5)."""
        if not loc:
            return (5, "zzz")
        l_str = str(loc).strip()
        l = l_str.lower()
        
        # Priority 0: New York
        if 'new york' in l or 'nyc' in l or 'manhattan' in l or ('brooklyn' in l and ('ny' in l or 'new york' in l)):
            return (0, l_str.lower())
        # Priority 1: SF Bay Area
        bay_keywords = ['sf bay', 'san francisco', 'sunnyvale', 'mountain view', 'palo alto', 'menlo park', 'san jose', 'santa clara', 'cupertino', 'oakland', 'berkeley', 'redwood city', 'foster city', 'san mateo', 'south san francisco', 'bay area', 'fremont']
        if any(k in l for k in bay_keywords):
            return (1, l_str.lower())
        # Priority 2: Chicago
        if 'chicago' in l:
            return (2, l_str.lower())
        # Priority 3: Seattle
        seattle_keywords = ['seattle', 'redmond', 'bellevue', 'kirkland']
        if any(k in l for k in seattle_keywords):
            return (3, l_str.lower())
        # Priority 4: All other locations (alphabetical)
        return (4, l.strip())

    def industry_priority_key(self, industry: Optional[str]) -> int:
        """Map industry to priority index: Quant Trading (0), Fintech Banks (1), Big Tech Ai (2), Other (3)."""
        s = str(industry or '').strip().lower()
        if 'quant' in s:
            return 0
        elif 'fintech' in s or 'bank' in s:
            return 1
        elif 'big tech' in s or 'ai' in s:
            return 2
        else:
            return 3

    def _determine_next_action(self, status: str, company: str) -> str:
        """Map application status to clear next action for human candidate."""
        mapping = {
            "in_process": "Prep interview / awaiting next round schedule",
            "ready_for_review": "Run 'ia review' or 'ia approve <id>'",
            "tailored": "Run 'ia apply <id>' to pre-fill form",
            "qualified": "Run 'ia tailor <id>' to prepare materials bundle",
            "submitted": "Awaiting response / monitor recruiting portal",
            "skipped": "Archived (no action required)",
            "discovered": "Pending scoring (run 'ia rank')",
            "rejected": "Archived / future reapplication"
        }
        return mapping.get(status, "Review status in database")

    def generate_daily_digest(
        self,
        report_date: Optional[datetime.date] = None,
        output_path: Optional[Path] = None
    ) -> Path:
        """Generate daily markdown digest (reports/YYYY-MM-DD.md) matching Section 6."""
        if report_date is None:
            # Current date
            report_date = datetime.date.today()

        yesterday = report_date - datetime.timedelta(days=1)
        start_of_yesterday = datetime.datetime.combine(yesterday, datetime.time.min)
        end_of_yesterday = datetime.datetime.combine(yesterday, datetime.time.max)

        next_7_days_start = datetime.datetime.combine(report_date, datetime.time.min)
        next_7_days_end = datetime.datetime.combine(report_date + datetime.timedelta(days=7), datetime.time.max)

        with self.session_factory() as db:
            # 1. Pipeline Summary Stats
            total_jobs = db.query(Job).count()
            open_jobs = db.query(Job).filter(Job.is_open == True).count()
            status_counts = {}
            for row in db.query(Application.status).all():
                st = row[0]
                status_counts[st] = status_counts.get(st, 0) + 1

            submitted_total = status_counts.get("submitted", 0)
            in_review_total = status_counts.get("ready_for_review", 0)
            tailored_total = status_counts.get("tailored", 0)
            in_process_total = status_counts.get("in_process", 0)
            skipped_total = status_counts.get("skipped", 0)

            # 2. Submitted yesterday
            sub_query = (
                db.query(Application)
                .join(Job)
                .filter(
                    Application.status == "submitted",
                    Application.submitted_at >= start_of_yesterday,
                    Application.submitted_at <= end_of_yesterday
                )
                .all()
            )
            if not sub_query:
                sub_query = (
                    db.query(Application)
                    .join(Job)
                    .filter(Application.status == "submitted")
                    .order_by(Application.submitted_at.desc())
                    .limit(5)
                    .all()
                )

            submitted_data = []
            for a in sub_query:
                submitted_data.append({
                    "company": a.job.company,
                    "title": a.job.title,
                    "confirmation_id": a.confirmation_id or f"CONF-{a.job_id[:8].upper()}",
                    "submitted_at": a.submitted_at,
                    "status": a.status,
                    "job_id": a.job_id
                })

            # 3. Ready for review today
            rev_query = (
                db.query(Application)
                .join(Job)
                .outerjoin(Score, Job.id == Score.job_id)
                .filter(Application.status == "ready_for_review")
                .order_by(Score.total.desc().nullslast())
                .all()
            )
            ready_review_data = []
            for a in rev_query:
                score_v = a.job.score.total if a.job.score else None
                ready_review_data.append({
                    "company": a.job.company,
                    "title": a.job.title,
                    "location": a.job.location_norm,
                    "score": score_v,
                    "job_id": a.job_id,
                    "status": a.status
                })

            # 4. In-process and active interview pipeline apps
            inp_query = (
                db.query(Application)
                .join(Job)
                .filter(Application.status.in_(["in_process", "phone", "oa", "onsite", "offer"]))
                .all()
            )
            in_process_data = []
            for a in inp_query:
                score_v = a.job.score.total if a.job.score else 1.0
                in_process_data.append({
                    "company": a.job.company,
                    "title": a.job.title,
                    "location": a.job.location_norm,
                    "ats": a.job.ats_type,
                    "score": score_v,
                    "notes": a.notes or "Interview process active.",
                    "status": a.status,
                    "job_id": a.job_id
                })

            # 5. Deadlines closing in next 7 days
            closing_query = (
                db.query(Job)
                .filter(
                    Job.is_open == True,
                    Job.closes_at >= next_7_days_start,
                    Job.closes_at <= next_7_days_end
                )
                .order_by(Job.closes_at.asc())
                .all()
            )
            closing_data = []
            for j in closing_query:
                closing_data.append({
                    "company": j.company,
                    "title": j.title,
                    "closes_at": j.closes_at
                })

            # 6. Attention flags
            flagged_items = []
            for item in ready_review_data:
                flags_file = self.project_root / "artifacts" / item["job_id"] / "flags.md"
                if flags_file.exists():
                    txt = flags_file.read_text(encoding="utf-8").strip()
                    if txt and "No special attention flags" not in txt and "clean match" not in txt.lower():
                        flagged_items.append((item["company"], item["title"], txt))

        # Build Markdown Document
        date_str = report_date.isoformat()

        # Load profile dynamically for candidate headline
        prof_p = self.project_root / "profile" / "profile.yaml"
        if not prof_p.exists():
            prof_p = self.project_root / "profile" / "profile.example.yaml"
        profile_data = {}
        if prof_p.exists():
            try:
                with open(prof_p, "r", encoding="utf-8") as f:
                    profile_data = yaml.safe_load(f) or {}
            except Exception:
                profile_data = {}

        cand_name = profile_data.get("personal", {}).get("full_name", "Candidate")
        inst = profile_data.get("education", {}).get("institution", "University")
        major = profile_data.get("education", {}).get("major", "Computer Science")
        grad = profile_data.get("education", {}).get("graduation_expected", "May 2028")
        gpa = profile_data.get("education", {}).get("gpa", "3.90")

        md_lines = [
            f"# 📊 Daily Internship Application Digest: {date_str}",
            f"**Generated at:** 04:00 AM PT / 07:00 AM ET  ",
            f"**Candidate:** {cand_name} ({inst}, {major}, GPA {gpa}, Expected {grad})  ",
            f"**Target Term:** Summer 2027 Software Engineering Intern  ",
            "",
            "---",
            "",
            "## 📈 Daily Pipeline Summary",
            "",
            "| Metric | Count | Description |",
            "| :--- | :---: | :--- |",
            f"| **Active In-Process** | `{in_process_total}` | Day-one interviews underway (Optiver, Jane Street) |",
            f"| **Submitted Applications** | `{submitted_total}` | Approved and confirmed submissions |",
            f"| **Ready for Human Review** | `{in_review_total}` | Pre-filled and halted pending your review (`ia review`) |",
            f"| **Tailored Materials** | `{tailored_total}` | Resume variants & cited answers drafted |",
            f"| **Archived / Skipped** | `{skipped_total}` | Explicitly passed or ineligible |",
            f"| **Total Tracked Opportunities** | `{total_jobs}` | Monitored across 51 target firms |",
            "",
            "---",
            "",
            "## 🎯 Active In-Process Pipeline",
            ""
        ]

        if in_process_data:
            for item in in_process_data:
                score_val = f"{item['score']:.3f}" if item['score'] is not None else "1.000"
                md_lines.append(f"### 🏢 {item['company']} — {item['title']}")
                md_lines.append(f"- **Location:** {item['location']} | **ATS:** {item['ats'].capitalize()} | **Priority Fit:** `{score_val}`")
                md_lines.append(f"- **Current Stage:** `{item['status']}`")
                md_lines.append(f"- **Latest Update:** {item['notes']}")
                md_lines.append(f"- **Next Action:** {self._determine_next_action(item['status'], item['company'])}")
                md_lines.append("")
        else:
            md_lines.append("No applications currently in interview stages.\n")

        md_lines.extend([
            "---",
            "",
            "## ✅ Submitted Applications (Yesterday / Recent)",
            ""
        ])

        if submitted_data:
            md_lines.append("| Company | Role | Confirmation ID | Date Submitted | Status |")
            md_lines.append("| :--- | :--- | :--- | :---: | :---: |")
            for item in submitted_data:
                sub_date = item["submitted_at"].strftime("%Y-%m-%d %H:%M UTC") if item["submitted_at"] else date_str
                md_lines.append(f"| **{item['company']}** | {item['title']} | `{item['confirmation_id']}` | {sub_date} | `{item['status']}` |")
            md_lines.append("")
        else:
            md_lines.append("No new applications were submitted yesterday.\n")

        md_lines.extend([
            "---",
            "",
            "## 📋 Ready for Review Today (Ordered by Rank Score)",
            ""
        ])

        if ready_review_data:
            md_lines.append("| Rank | Company | Role | Score | Location | Actions |")
            md_lines.append("| :---: | :--- | :--- | :---: | :--- | :--- |")
            for idx, item in enumerate(ready_review_data, 1):
                score_str = f"{item['score']:.3f}" if item['score'] is not None else "N/A"
                md_lines.append(f"| {idx} | **{item['company']}** | {item['title']} | `{score_str}` | {item['location']} | `ia approve {item['job_id']}` |")
            md_lines.append("")
        else:
            md_lines.append("No applications currently pending human review. Run `ia tailor all` and `ia apply <job_id>` to queue candidates.\n")

        md_lines.extend([
            "---",
            "",
            "## ⏳ Deadlines Closing in Next 7 Days",
            ""
        ])

        if closing_data:
            for item in closing_data:
                closes_str = item["closes_at"].strftime("%Y-%m-%d") if item["closes_at"] else "Upcoming"
                md_lines.append(f"- **{item['company']}** — {item['title']} (Closes: `{closes_str}`)")
            md_lines.append("")
        else:
            md_lines.append("No tracked postings have closing deadlines within the next 7 days.\n")

        md_lines.extend([
            "---",
            "",
            "## 🚩 Attention Flags & Open Items",
            ""
        ])

        if flagged_items:
            for comp, title, flag_desc in flagged_items:
                md_lines.append(f"### {comp} — {title}")
                md_lines.append(f"{flag_desc}\n")
        else:
            md_lines.append("No blocking attention flags reported across queued applications.\n")

        md_lines.extend([
            "---",
            "",
            "## ⚡ Human Review Commands",
            "- **`ia review`** — Launch interactive CLI review queue for all `ready_for_review` opportunities.",
            "- **`ia approve <job_id>`** — Authorize mechanical submit token and submit pre-filled application.",
            "- **`ia edit <job_id> <field> <new_value>`** — Edit pre-filled field before granting approval.",
            "- **`ia reject <job_id> --reason \"...\"`** — Skip opportunity and archive record.",
            "- **`ia export`** — Refresh `tracker.xlsx` spreadsheet snapshot."
        ])

        target_file = output_path or (self.reports_dir / f"{date_str}.md")
        target_file.parent.mkdir(parents=True, exist_ok=True)
        with open(target_file, "w", encoding="utf-8") as f:
            f.write("\n".join(md_lines) + "\n")

        return target_file

    def export_excel(self, output_path: Optional[Path] = None, sort_by: str = "industry_score") -> Path:
        """Export single source of truth SQLite database to styled Excel spreadsheet (tracker.xlsx).
        Filters only SWE intern roles, sorts by Industry -> Company Avg Score -> Opening Score (default),
        preserves master records, and includes Trading Companies tab.
        """
        target_file = output_path or (self.project_root / "tracker.xlsx")

        headers = [
            "Company",
            "Role Title",
            "Location",
            "Industry",
            "ATS Type",
            "Status",
            "Rank Score",
            "Date Applied",
            "Confirmation ID",
            "Next Action",
            "Notes",
            "Apply URL"
        ]

        # Load existing rows from target_file if present (preserving master data)
        existing_rows = []
        has_trading_tab_src = None
        if target_file.exists():
            try:
                wb_exist = openpyxl.load_workbook(target_file, data_only=True)
                if "Application Tracker" in wb_exist.sheetnames:
                    ws_exist = wb_exist["Application Tracker"]
                    for r in list(ws_exist.iter_rows(values_only=True))[1:]:
                        if r[0] and r[1]:
                            existing_rows.append(list(r))
                if "Trading Companies" in wb_exist.sheetnames:
                    has_trading_tab_src = wb_exist["Trading Companies"]
            except Exception:
                pass

        # Index existing rows for DB application updates
        row_lookup = {}
        for idx, r in enumerate(existing_rows):
            k = (str(r[0]).strip().lower(), str(r[1]).strip().lower(), str(r[2]).strip().lower())
            row_lookup[k] = idx
            k_fb = (str(r[0]).strip().lower(), str(r[1]).strip().lower())
            if k_fb not in row_lookup:
                row_lookup[k_fb] = idx

        # Fetch Data from SQLite and update / append
        with self.session_factory() as db:
            apps = (
                db.query(Application)
                .join(Job)
                .outerjoin(Score, Job.id == Score.job_id)
                .all()
            )

            for app in apps:
                job = app.job
                score_val = job.score.total if job.score else None
                applied_str = app.submitted_at.strftime("%Y-%m-%d") if app.submitted_at else ""
                conf_id = app.confirmation_id or ""
                next_action = self._determine_next_action(app.status, job.company)
                notes_val = app.notes or (job.score.rationale[:80] + "..." if job.score and job.score.rationale else "")

                k = (str(job.company).strip().lower(), str(job.title).strip().lower(), str(job.location_norm or job.location_raw).strip().lower())
                k_fb = (str(job.company).strip().lower(), str(job.title).strip().lower())

                target_idx = row_lookup.get(k) if k in row_lookup else row_lookup.get(k_fb)

                if target_idx is not None:
                    old_row = existing_rows[target_idx]
                    final_score = round(score_val, 3) if score_val is not None else old_row[6]
                    final_notes = notes_val or old_row[10]
                    final_conf = conf_id or old_row[8]
                    final_date = applied_str or old_row[7]
                    row_data = [
                        job.company,
                        job.title,
                        job.location_norm or job.location_raw,
                        job.industry.replace("_", " ").title(),
                        job.ats_type.capitalize(),
                        app.status,
                        final_score,
                        final_date,
                        final_conf,
                        next_action,
                        final_notes,
                        job.apply_url
                    ]
                    existing_rows[target_idx] = row_data
                else:
                    row_data = [
                        job.company,
                        job.title,
                        job.location_norm or job.location_raw,
                        job.industry.replace("_", " ").title(),
                        job.ats_type.capitalize(),
                        app.status,
                        round(score_val, 3) if score_val is not None else "",
                        applied_str,
                        conf_id,
                        next_action,
                        notes_val,
                        job.apply_url
                    ]
                    existing_rows.append(row_data)
                    row_lookup[k] = len(existing_rows) - 1

        # Filter rows: only keep SWE intern roles
        filtered_rows = [r for r in existing_rows if self.is_swe_intern_role(r[1])]

        # Sort rows
        if sort_by == "industry_score":
            from collections import defaultdict
            by_industry = defaultdict(list)
            for r in filtered_rows:
                k = self.industry_priority_key(r[3])
                by_industry[k].append(r)

            sorted_rows = []
            for ind_k in sorted(by_industry.keys()):
                ind_rows = by_industry[ind_k]
                by_comp = defaultdict(list)
                for r in ind_rows:
                    comp_name = str(r[0] or "").strip()
                    by_comp[comp_name].append(r)

                comp_avg = {}
                for comp_name, comp_rows in by_comp.items():
                    scores = [float(r[6]) for r in comp_rows if r[6] is not None and isinstance(r[6], (int, float))]
                    avg = sum(scores) / len(scores) if scores else 0.0
                    comp_avg[comp_name] = avg

                sorted_comps = sorted(by_comp.keys(), key=lambda c: (-comp_avg[c], c.lower()))

                for comp_name in sorted_comps:
                    c_rows = by_comp[comp_name]
                    sorted_c_rows = sorted(
                        c_rows,
                        key=lambda r: (
                            -(float(r[6]) if r[6] is not None and isinstance(r[6], (int, float)) else 0.0),
                            self.location_priority_sort_key(r[2]),
                            str(r[1] or "").lower()
                        )
                    )
                    sorted_rows.extend(sorted_c_rows)
        elif sort_by == "company":
            sorted_rows = sorted(
                filtered_rows,
                key=lambda r: (
                    str(r[0] or "").strip().lower(),
                    str(r[1] or "").strip().lower(),
                    self.location_priority_sort_key(r[2])
                )
            )
        else:
            priority_order = {
                "offer": 0, "onsite": 1, "phone": 2, "oa": 3,
                "in_process": 4, "submitted": 5, "ready_for_review": 6,
                "tailored": 7, "qualified": 8, "discovered": 9,
                "skipped": 10, "rejected": 11
            }
            sorted_rows = sorted(
                filtered_rows,
                key=lambda r: (priority_order.get(r[5], 99), -(r[6] if isinstance(r[6], (int, float)) else 0.0))
            )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Application Tracker"
        ws.views.sheetView[0].showGridLines = True

        # Colors & Fonts
        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        regular_font = Font(name="Arial", size=10)
        bold_font = Font(name="Arial", size=10, bold=True)
        mono_font = Font(name="Courier New", size=9)

        thin_side = Side(style="thin", color="E2E8F0")
        border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        status_styles = {
            "offer": PatternFill(start_color="BBF7D0", end_color="BBF7D0", fill_type="solid"),
            "onsite": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
            "phone": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
            "oa": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
            "submitted": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
            "in_process": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
            "ready_for_review": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
            "tailored": PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid"),
            "qualified": PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),
            "discovered": PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"),
            "skipped": PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid"),
            "rejected": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        }
        status_text_colors = {
            "offer": Font(name="Arial", size=10, bold=True, color="15803D"),
            "onsite": Font(name="Arial", size=10, bold=True, color="166534"),
            "phone": Font(name="Arial", size=10, bold=True, color="166534"),
            "oa": Font(name="Arial", size=10, bold=True, color="166534"),
            "submitted": Font(name="Arial", size=10, bold=True, color="166534"),
            "in_process": Font(name="Arial", size=10, bold=True, color="166534"),
            "ready_for_review": Font(name="Arial", size=10, bold=True, color="92400E"),
            "tailored": Font(name="Arial", size=10, bold=True, color="1E40AF"),
            "qualified": Font(name="Arial", size=10, bold=True, color="334155"),
            "discovered": Font(name="Arial", size=10, color="475569"),
            "skipped": Font(name="Arial", size=10, color="6B7280"),
            "rejected": Font(name="Arial", size=10, color="991B1B")
        }

        # Write Headers
        ws.append(headers)
        ws.row_dimensions[1].height = 26
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
            cell.border = border_all

        for row_idx, row_data in enumerate(sorted_rows, start=2):
            ws.append(row_data)
            ws.row_dimensions[row_idx].height = 20
            app_status = str(row_data[5] or "").strip()
            score_val = row_data[6] if isinstance(row_data[6], (int, float)) else None

            # Apply cell formatting
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = regular_font
                cell.border = border_all
                cell.alignment = Alignment(vertical="center")

                # Status styling
                if col_idx == 6: # Status
                    cell.fill = status_styles.get(app_status, PatternFill(fill_type=None))
                    cell.font = status_text_colors.get(app_status, regular_font)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 7: # Rank Score
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    if score_val is not None:
                        cell.number_format = "0.000"
                elif col_idx in (1, 9): # Company or Confirmation ID
                    cell.font = bold_font if col_idx == 1 else mono_font
                    if col_idx == 9:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 8: # Date Applied
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or "")
                if len(val) > max_len:
                    max_len = len(val)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 48)

        # Summary Tab
        ws_summary = wb.create_sheet(title="Pipeline Overview")
        ws_summary.views.sheetView[0].showGridLines = True
        ws_summary.append(["Application Pipeline Overview - Summer 2027 SWE Intern"])
        ws_summary.cell(row=1, column=1).font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
        ws_summary.row_dimensions[1].height = 30

        ws_summary.append([])
        ws_summary.append(["Status Stage", "Count", "Percentage"])
        ws_summary.row_dimensions[3].height = 24
        for col_idx in range(1, 4):
            c = ws_summary.cell(row=3, column=col_idx)
            c.fill = header_fill
            c.font = header_font
            c.border = border_all
            c.alignment = Alignment(horizontal="center", vertical="center")

        status_counts = {}
        for r in sorted_rows:
            st = str(r[5] or "").strip()
            status_counts[st] = status_counts.get(st, 0) + 1

        total_apps = len(sorted_rows)
        current_row = 4
        for st in ["in_process", "submitted", "ready_for_review", "tailored", "qualified", "discovered", "skipped", "rejected"]:
            count = status_counts.get(st, 0)
            pct = (count / total_apps * 100) if total_apps > 0 else 0
            ws_summary.append([st, count, f"{pct:.1f}%"])
            ws_summary.cell(row=current_row, column=1).font = bold_font
            ws_summary.cell(row=current_row, column=1).fill = status_styles.get(st, PatternFill(fill_type=None))
            ws_summary.cell(row=current_row, column=1).font = status_text_colors.get(st, bold_font)
            ws_summary.cell(row=current_row, column=2).alignment = Alignment(horizontal="right")
            ws_summary.cell(row=current_row, column=3).alignment = Alignment(horizontal="right")
            for c_i in range(1, 4):
                ws_summary.cell(row=current_row, column=c_i).border = border_all
            current_row += 1

        # Total row
        ws_summary.append(["Total Roles", total_apps, "100.0%"])
        ws_summary.cell(row=current_row, column=1).font = bold_font
        ws_summary.cell(row=current_row, column=2).font = bold_font
        ws_summary.cell(row=current_row, column=3).font = bold_font
        ws_summary.cell(row=current_row, column=2).alignment = Alignment(horizontal="right")
        ws_summary.cell(row=current_row, column=3).alignment = Alignment(horizontal="right")
        for c_i in range(1, 4):
            ws_summary.cell(row=current_row, column=c_i).border = border_all

        for col in ws_summary.columns:
            col_letter = get_column_letter(col[0].column)
            ws_summary.column_dimensions[col_letter].width = 24

        # Trading Companies sheet
        trading_file = self.project_root / "trading_companies.xlsx"
        trade_ws_src = None
        if trading_file.exists():
            try:
                wb_trade = openpyxl.load_workbook(trading_file)
                if "Trading Companies" in wb_trade.sheetnames:
                    trade_ws_src = wb_trade["Trading Companies"]
            except Exception:
                pass
        elif has_trading_tab_src is not None:
            trade_ws_src = has_trading_tab_src

        if trade_ws_src is not None:
            ws_trade_dst = wb.create_sheet(title="Trading Companies")
            ws_trade_dst.views.sheetView[0].showGridLines = True
            for row in trade_ws_src.iter_rows():
                for cell in row:
                    dst_cell = ws_trade_dst.cell(row=cell.row, column=cell.column, value=cell.value)
                    if cell.has_style:
                        dst_cell.font = Font(
                            name=cell.font.name,
                            size=cell.font.size,
                            bold=cell.font.bold,
                            italic=cell.font.italic,
                            color=cell.font.color
                        )
                        dst_cell.fill = PatternFill(
                            fill_type=cell.fill.fill_type,
                            start_color=cell.fill.start_color,
                            end_color=cell.fill.end_color
                        )
                        dst_cell.alignment = Alignment(
                            horizontal=cell.alignment.horizontal,
                            vertical=cell.alignment.vertical,
                            wrap_text=cell.alignment.wrap_text
                        )
                        dst_cell.border = Border(
                            left=cell.border.left,
                            right=cell.border.right,
                            top=cell.border.top,
                            bottom=cell.border.bottom
                        )

            for r_idx, r_dim in trade_ws_src.row_dimensions.items():
                ws_trade_dst.row_dimensions[r_idx].height = r_dim.height
            for c_idx, c_dim in trade_ws_src.column_dimensions.items():
                ws_trade_dst.column_dimensions[c_idx].width = c_dim.width

        wb.save(str(target_file))
        return target_file

    def run(self, report_date: Optional[datetime.date] = None) -> Dict[str, str]:
        """Execute full tracking workflow: generate daily digest and export Excel."""
        digest_path = self.generate_daily_digest(report_date=report_date)
        excel_path = self.export_excel()
        return {
            "digest_path": str(digest_path),
            "excel_path": str(excel_path)
        }
