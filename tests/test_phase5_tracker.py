import os
import json
import pytest
import datetime
from pathlib import Path
from typer.testing import CliRunner
import openpyxl

from db.session import init_db, SessionLocal, generate_id
from db.models import Job, Application, Score, Artifact, Event
from agents.rank import FilterAndRankAgent
from agents.tailor import TailoringAgent
from agents.apply import ApplicationAgent
from agents.tracker import TrackerAgent
from cli import app

PROJECT_ROOT = Path(__file__).resolve().parent.parent
FIXTURE_HTML = PROJECT_ROOT / "tests" / "fixtures" / "greenhouse_form.html"
runner = CliRunner()

@pytest.fixture(autouse=True)
def setup_db():
    init_db()

def test_end_to_end_3_synthetic_postings_flow():
    """Phase 5 Acceptance Test:
    Run an end-to-end flow on 3 synthetic postings through:
    discovery -> rank -> tailor -> pre-fill -> approve & submit (for 2) + 1 in review.
    Run TrackerAgent to generate daily digest markdown and Excel tracker.
    Verify digest accurately reflects all state changes and Excel matches DB state.
    """
    now = datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None)
    synthetic_specs = [
        {
            "id": "job_synth_p5_1",
            "company": "Citadel",
            "title": "Software Engineer Intern - Core Infrastructure (Summer 2027)",
            "industry": "quant_trading",
            "location_raw": "Chicago, IL",
            "location_norm": "Chicago, IL",
            "ats": "greenhouse",
            "should_submit": True
        },
        {
            "id": "job_synth_p5_2",
            "company": "Anthropic",
            "title": "Software Engineer Intern - Systems & Compute (Summer 2027)",
            "industry": "big_tech_ai",
            "location_raw": "San Francisco, CA",
            "location_norm": "SF Bay Area",
            "ats": "greenhouse",
            "should_submit": True
        },
        {
            "id": "job_synth_p5_3",
            "company": "Stripe",
            "title": "Software Engineering Intern - Platform & APIs (Summer 2027)",
            "industry": "fintech_banks",
            "location_raw": "Seattle, WA",
            "location_norm": "Seattle, WA",
            "ats": "greenhouse",
            "should_submit": False # Stays in ready_for_review
        }
    ]

    with SessionLocal() as db:
        for spec in synthetic_specs:
            jid = spec["id"]
            # Clean if exists
            existing_app = db.query(Application).filter_by(job_id=jid).first()
            if existing_app:
                db.query(Event).filter_by(application_id=existing_app.id).delete(synchronize_session=False)
                db.query(Artifact).filter_by(application_id=existing_app.id).delete(synchronize_session=False)
                db.delete(existing_app)
            db.query(Score).filter_by(job_id=jid).delete(synchronize_session=False)
            existing_job = db.query(Job).filter_by(id=jid).first()
            if existing_job:
                db.delete(existing_job)
            db.commit()

            # 1. Discovery stage
            job = Job(
                id=jid,
                dedup_key=f"dedup_{jid}",
                company=spec["company"],
                title=spec["title"],
                location_raw=spec["location_raw"],
                location_norm=spec["location_norm"],
                industry=spec["industry"],
                ats_type=spec["ats"],
                url=f"file://{FIXTURE_HTML.resolve()}",
                apply_url=f"file://{FIXTURE_HTML.resolve()}",
                source="synthetic_test",
                first_seen=now,
                last_seen=now,
                is_open=True
            )
            db.add(job)

            app_rec = Application(
                id=f"app_{jid}",
                job_id=jid,
                status="discovered",
                last_event_at=now
            )
            db.add(app_rec)
            db.commit()

    # 2. Rank stage
    ranker = FilterAndRankAgent()
    for spec in synthetic_specs:
        score_data = ranker.rank_job(spec["id"])
        assert score_data["eligible"] is True
        assert score_data["total"] >= 0.40

    # 3. Tailor stage
    tailor = TailoringAgent()
    for spec in synthetic_specs:
        bundle = tailor.tailor_job(spec["id"])
        assert Path(bundle["resume"]).exists()
        assert Path(bundle["answers"]).exists()

    # 4. Pre-fill stage (Playwright automation)
    app_agent = ApplicationAgent()
    import asyncio
    for spec in synthetic_specs:
        res = asyncio.run(app_agent.prefill_application(spec["id"], custom_page_url=f"file://{FIXTURE_HTML.resolve()}"))
        assert res["status"] == "ready_for_review"

    # 5. Approval and Submission stage
    for spec in synthetic_specs:
        if spec["should_submit"]:
            sub_res = asyncio.run(app_agent.approve_and_submit(spec["id"], actor="human", custom_page_url=f"file://{FIXTURE_HTML.resolve()}"))
            assert sub_res["status"] == "submitted"
            assert sub_res["confirmation_id"].startswith("HRT-") or sub_res["confirmation_id"].startswith("CONF-")

    # 6. Run Tracker Agent
    tracker = TrackerAgent()
    out = tracker.run()
    digest_path = Path(out["digest_path"])
    excel_path = Path(out["excel_path"])

    assert digest_path.exists()
    assert excel_path.exists()

    # Verify Daily Digest Contents
    digest_text = digest_path.read_text(encoding="utf-8")
    assert "Daily Internship Application Digest" in digest_text
    # In-process pipeline preserved
    assert "Optiver" in digest_text
    assert "Jane Street" in digest_text
    # Submitted applications
    assert "Citadel" in digest_text
    assert "Anthropic" in digest_text
    # Application ready for review
    assert "Stripe" in digest_text
    assert "ia approve job_synth_p5_3" in digest_text

    # Verify Excel Spreadsheet Matches DB State
    wb = openpyxl.load_workbook(excel_path)
    assert "Application Tracker" in wb.sheetnames
    assert "Pipeline Overview" in wb.sheetnames

    ws = wb["Application Tracker"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    assert "Company" in header
    assert "Status" in header
    assert "Rank Score" in header
    assert "Confirmation ID" in header

    # Build role -> row dict to avoid company name collisions with existing jobs
    role_rows = {}
    for r in rows[1:]:
        role_title = r[1]
        if role_title:
            role_rows[role_title] = r

    # Verify Citadel row
    citadel_row = role_rows.get("Software Engineer Intern - Core Infrastructure (Summer 2027)")
    assert citadel_row is not None
    assert citadel_row[0] == "Citadel"
    assert citadel_row[5] == "submitted"
    assert citadel_row[8] != "" # Confirmation ID present

    # Verify Anthropic row
    anthropic_row = role_rows.get("Software Engineer Intern - Systems & Compute (Summer 2027)")
    assert anthropic_row is not None
    assert anthropic_row[0] == "Anthropic"
    assert anthropic_row[5] == "submitted"

    # Verify Stripe row
    stripe_row = role_rows.get("Software Engineering Intern - Platform & APIs (Summer 2027)")
    assert stripe_row is not None
    assert stripe_row[0] == "Stripe"
    assert stripe_row[5] == "ready_for_review"
    assert "ia review" in str(stripe_row[9]) or "ia approve" in str(stripe_row[9])

    # Verify Jane Street and Optiver exist
    companies = [r[0] for r in rows[1:]]
    assert "Jane Street" in companies
    assert "Optiver" in companies

    # Teardown synthetic records so they do not pollute live database
    with SessionLocal() as db:
        synth_ids = [s["id"] for s in synthetic_specs]
        app_ids = [f"app_{s['id']}" for s in synthetic_specs]
        db.query(Artifact).filter(Artifact.application_id.in_(app_ids)).delete(synchronize_session=False)
        db.query(Score).filter(Score.job_id.in_(synth_ids)).delete(synchronize_session=False)
        db.query(Event).filter(Event.application_id.in_(app_ids)).delete(synchronize_session=False)
        db.query(Application).filter(Application.id.in_(app_ids)).delete(synchronize_session=False)
        db.query(Job).filter(Job.id.in_(synth_ids)).delete(synchronize_session=False)
        db.commit()

def test_digest_across_3_consecutive_simulated_dates():
    """Verify daily digest reflects state changes accurately across 3 consecutive simulated dates."""
    tracker = TrackerAgent()
    day1 = datetime.date(2026, 9, 18)
    day2 = datetime.date(2026, 9, 19)
    day3 = datetime.date(2026, 9, 20)

    p1 = tracker.generate_daily_digest(report_date=day1)
    p2 = tracker.generate_daily_digest(report_date=day2)
    p3 = tracker.generate_daily_digest(report_date=day3)

    assert p1.name == "2026-09-18.md"
    assert p2.name == "2026-09-19.md"
    assert p3.name == "2026-09-20.md"

    t1 = p1.read_text(encoding="utf-8")
    t2 = p2.read_text(encoding="utf-8")
    t3 = p3.read_text(encoding="utf-8")

    assert "Daily Internship Application Digest: 2026-09-18" in t1
    assert "Daily Internship Application Digest: 2026-09-19" in t2
    assert "Daily Internship Application Digest: 2026-09-20" in t3

    # All 3 digests accurately track the active interview pipeline
    for t in [t1, t2, t3]:
        assert "Optiver" in t
        assert "Jane Street" in t
        assert "phone" in t.lower()

def test_excel_styling_and_conditional_formatting():
    """Verify Excel workbook styles, headers, colors, and layout."""
    tracker = TrackerAgent()
    excel_path = tracker.export_excel()
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["Application Tracker"]

    # Header styling
    header_cell = ws.cell(row=1, column=1)
    assert header_cell.fill.start_color.rgb == "001E3A8A" or "1E3A8A" in str(header_cell.fill.start_color.rgb)
    assert header_cell.font.bold is True
    assert header_cell.font.color.rgb == "00FFFFFF" or "FFFFFF" in str(header_cell.font.color.rgb)

    # Status column conditional formatting
    status_col_idx = 6
    found_submitted = False
    found_in_process = False

    for row in range(2, min(ws.max_row + 1, 30)):
        val = ws.cell(row=row, column=status_col_idx).value
        cell_fill = ws.cell(row=row, column=status_col_idx).fill
        if val in ["submitted", "in_process"]:
            # Green fill
            fill_hex = str(cell_fill.start_color.rgb)
            assert "DCFCE7" in fill_hex or "00DCFCE7" in fill_hex
            if val == "submitted":
                found_submitted = True
            if val == "in_process":
                found_in_process = True

    assert found_in_process, "In-process rows must have green fill"

    # Overview sheet verification
    ws_overview = wb["Pipeline Overview"]
    assert ws_overview.max_row >= 7
    title_cell = ws_overview.cell(row=1, column=1).value
    assert "Pipeline Overview" in title_cell

def test_cli_digest_and_export_commands():
    """Verify ia digest and ia export CLI commands."""
    # Test ia digest with default date
    res_digest = runner.invoke(app, ["digest"])
    assert res_digest.exit_code == 0
    assert "Daily digest report generated" in res_digest.output
    assert "Excel spreadsheet updated" in res_digest.output

    # Test ia digest with explicit date
    res_date = runner.invoke(app, ["digest", "--date", "2026-09-21"])
    assert res_date.exit_code == 0
    assert (PROJECT_ROOT / "reports" / "2026-09-21.md").exists()

    # Test ia export
    res_export = runner.invoke(app, ["export"])
    assert res_export.exit_code == 0
    assert "Pipeline exported successfully" in res_export.output
